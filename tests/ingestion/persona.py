"""The fake persona every ingestion test uses. Nothing here is a real person, company or ID."""

from __future__ import annotations

import io
import zipfile

NAME = "Kavya Testpal"
PUBLIC_EMAIL = "kavya.testpal@example.com"
PRIVATE_EMAIL = "kavya.private@example.org"
COMPANY = "Acme Widgets"
PAN = "ABCPT1234K"
IFSC = "ABCD0123456"
BANK_ACCOUNT = "123456789012345"
PASSPORT = "Z1234567"
UAN = "100200300400"
DOB = "14/03/1995"
ROLL_NUMBER = "2013CS0420"


def note(*extra: str) -> str:
    return "\n".join(
        [
            "# Interview at Acme Widgets",
            "",
            f"{NAME} interviewed for a backend role. Reach me at {PUBLIC_EMAIL}.",
            "## Round 1",
            "System design of a URL shortener with consistent hashing.",
            *extra,
        ]
    )


def pdf_bytes(text: str) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=12)
    return bytes(doc.tobytes())


def image_pdf_bytes() -> bytes:
    """A one-page PDF whose only content is an image: no text layer, so OCR is needed."""
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(pymupdf.Rect(72, 72, 272, 172), stream=png_bytes())
    return bytes(doc.tobytes())


def png_bytes() -> bytes:
    import pymupdf

    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 20), False)
    pix.clear_with(255)
    return bytes(pix.tobytes("png"))


def docx_bytes(paragraph: str) -> bytes:
    parts = {
        "[Content_Types].xml": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" '
            'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/'
            'vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
        ),
        "_rels/.rels": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/'
            '2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>'
        ),
        "word/document.xml": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f"<w:body><w:p><w:r><w:t>{paragraph}</w:t></w:r></w:p></w:body></w:document>"
        ),
    }
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, body in parts.items():
            archive.writestr(name, body)
    return buffer.getvalue()


def xlsx_bytes(rows: list[list[str]]) -> bytes:
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def pptx_bytes(title: str) -> bytes:
    import pptx

    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = title
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()
